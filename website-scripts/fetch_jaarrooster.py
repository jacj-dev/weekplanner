"""
Zet data/jaarrooster.xlsx om naar data/jaarrooster.json, zodat de
weekplanner bij elke les kan laten zien wat er volgens de jaarplanning
gepland staat.

Opbouw van het Excel-bestand (per tabblad):
  - rij 1 = kopjes: "Dag", "Datum", "Bijzonderheden", en daarna één kolom
    per klas/vak (bv. "1BK - wi")
  - elke volgende rij = een week (alleen maandag) of een losse les
    (ma/wo/do); allebei mag door elkaar
  - een rij zonder datum hoort bij de laatste datum erboven

Welke Excel-kolom bij welke les in het Zermelo-rooster hoort, staat in
KOPPELINGEN hieronder. Een roostercode ziet eruit als
"v219 wi lv1d" of "e013 bi vs10k.bi1 vs10m.bi1": lokaal, vak, klas(sen).
Een koppeling ("wi", "lv1d") betekent: vak is "wi" en één van de klassen
begint met "lv1d".

Komt er een klas bij of verandert een kolomnaam? Pas dan alleen
KOPPELINGEN aan.

Draaien: python website-scripts/fetch_jaarrooster.py
"""

import datetime as dt
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
INVOER = ROOT / "data" / "jaarrooster.xlsx"
UITVOER = ROOT / "data" / "jaarrooster.json"

# Excel-kolomkop -> lijst van (vak, klas-begin) uit het Zermelo-rooster
KOPPELINGEN = {
    "1BK - wi": [("wi", "lv1a")],
    "1KGT": [("wi", "lv1d")],
    "Rekenen": [("rek", "lv1a"), ("rek", "lv1d")],
    "1BK - bi": [("bi", "lv1a")],
    "LPV2 - wi": [("wi", "lpv2")],
    "LPV2 - bi": [("bi", "lpv2")],
    "Vs9": [("bi", "vs9")],
    "Vs10": [("bi", "vs10")],
}

VASTE_KOPPEN = {"dag", "datum", "bijzonderheden"}


def naar_datum(waarde):
    if isinstance(waarde, dt.datetime):
        return waarde.date()
    if isinstance(waarde, dt.date):
        return waarde
    if isinstance(waarde, str):
        try:
            return dt.date.fromisoformat(waarde.strip()[:10])
        except ValueError:
            return None
    return None


def schoon(tekst):
    if tekst is None:
        return ""
    return " ".join(str(tekst).split())


def lees_tabblad(ws, planning, bijzonderheden, onbekend):
    koppen = [schoon(c.value) for c in ws[1]]
    try:
        datum_kol = [k.lower() for k in koppen].index("datum")
    except ValueError:
        print(f"Tabblad '{ws.title}' heeft geen kolom 'Datum', overgeslagen.")
        return
    bijz_kol = next((i for i, k in enumerate(koppen) if k.lower() == "bijzonderheden"), None)

    klaskolommen = {}
    for i, kop in enumerate(koppen):
        if not kop or kop.lower() in VASTE_KOPPEN:
            continue
        if kop in KOPPELINGEN:
            klaskolommen[i] = kop
        else:
            onbekend.add(f"{ws.title}: '{kop}'")

    laatste_datum = None
    for rij in ws.iter_rows(min_row=2, values_only=True):
        datum = naar_datum(rij[datum_kol]) if datum_kol < len(rij) else None
        heeft_eigen_datum = datum is not None
        if datum is None:
            datum = laatste_datum
        else:
            laatste_datum = datum
        if datum is None:
            continue
        maandag = datum - dt.timedelta(days=datum.weekday())
        week = maandag.isoformat()

        if bijz_kol is not None and bijz_kol < len(rij):
            tekst = schoon(rij[bijz_kol])
            if tekst and tekst not in bijzonderheden.setdefault(week, []):
                bijzonderheden[week].append(tekst)

        for i, kop in klaskolommen.items():
            if i >= len(rij):
                continue
            tekst = schoon(rij[i])
            if not tekst:
                continue
            planning.setdefault(kop, {}).setdefault(week, []).append(
                {"datum": datum.isoformat() if heeft_eigen_datum else None, "tekst": tekst}
            )


def main():
    if not INVOER.exists():
        print(f"{INVOER} niet gevonden.")
        sys.exit(1)

    wb = openpyxl.load_workbook(INVOER, data_only=True, read_only=True)
    planning, bijzonderheden, onbekend = {}, {}, set()
    for ws in wb.worksheets:
        lees_tabblad(ws, planning, bijzonderheden, onbekend)

    if onbekend:
        print("Let op, deze kolommen staan niet in KOPPELINGEN en zijn overgeslagen:")
        for k in sorted(onbekend):
            print("  -", k)

    uit = {
        "koppelingen": [
            {"kolom": kop, "match": [{"vak": v, "klas": k} for v, k in paren]}
            for kop, paren in KOPPELINGEN.items()
        ],
        "planning": planning,
        "bijzonderheden": bijzonderheden,
    }
    UITVOER.write_text(json.dumps(uit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    regels = sum(len(r) for weken in planning.values() for r in weken.values())
    print(f"{UITVOER.name} geschreven: {len(planning)} kolommen, {regels} planningsregels.")


if __name__ == "__main__":
    main()
